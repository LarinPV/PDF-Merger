"""
Unit and integration tests for PDF Merger.
Tests bookmarks preservation, image conversion, DOCX conversion, and merged outputs.
"""

import os
import sys
import tempfile
import unittest

# Ensure src module is in path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pypdf
from PIL import Image, ImageDraw
import docx

from src.converter import convert_image_to_pdf, convert_docx_fallback, convert_text_to_pdf
from src.merger import merge_files


class TestPDFMerger(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="pdf_test_")
        
    def tearDown(self):
        try:
            for root, dirs, files in os.walk(self.test_dir, topdown=False):
                for f in files:
                    os.remove(os.path.join(root, f))
                for d in dirs:
                    os.rmdir(os.path.join(root, d))
            os.rmdir(self.test_dir)
        except Exception:
            pass

    def test_pdf_with_bookmarks_merge(self):
        # 1. Create PDF 1 with bookmarks: 'Chapter 1' (p0), 'Chapter 2' (p1)
        pdf1_path = os.path.join(self.test_dir, "doc1.pdf")
        w1 = pypdf.PdfWriter()
        w1.add_blank_page(width=400, height=400)
        w1.add_outline_item("Глава 1: Введение", 0)
        w1.add_blank_page(width=400, height=400)
        w1.add_outline_item("Глава 2: Описание", 1)
        with open(pdf1_path, "wb") as f:
            w1.write(f)

        # 2. Create PDF 2 with bookmark: 'Раздел А' (p0)
        pdf2_path = os.path.join(self.test_dir, "doc2.pdf")
        w2 = pypdf.PdfWriter()
        w2.add_blank_page(width=400, height=400)
        w2.add_outline_item("Раздел А: Данные", 0)
        with open(pdf2_path, "wb") as f:
            w2.write(f)

        # 3. Create sample Image
        img_path = os.path.join(self.test_dir, "photo.jpg")
        img = Image.new("RGB", (300, 200), color=(100, 150, 200))
        d = ImageDraw.Draw(img)
        d.text((20, 20), "Test Image", fill=(255, 255, 255))
        img.save(img_path)

        # 4. Create sample DOCX
        docx_path = os.path.join(self.test_dir, "report.docx")
        doc = docx.Document()
        doc.add_heading("Заголовок Отчета", 0)
        doc.add_paragraph("Тестовый абзац документа Word.")
        doc.save(docx_path)

        # 5. Merge all files
        merged_out = os.path.join(self.test_dir, "combined.pdf")
        result = merge_files(
            file_list=[pdf1_path, pdf2_path, img_path, docx_path],
            output_path=merged_out,
            preserve_bookmarks=True,
            add_file_bookmarks=True
        )

        self.assertTrue(result["success"])
        self.assertTrue(os.path.exists(merged_out))
        self.assertGreater(os.path.getsize(merged_out), 1000)

        # Verify pages
        reader = pypdf.PdfReader(merged_out)
        # doc1 (2 pages) + doc2 (1 page) + photo (1 page) + docx (1 page) = 5 pages
        self.assertEqual(len(reader.pages), 5)

        # Verify outlines / bookmarks
        outline = reader.outline
        self.assertTrue(len(outline) > 0)
        
        # Flatten outline titles to check preservation
        def extract_titles(outline_list):
            titles = []
            for item in outline_list:
                if isinstance(item, list):
                    titles.extend(extract_titles(item))
                elif isinstance(item, dict) or hasattr(item, 'title'):
                    title = item.title if hasattr(item, 'title') else item.get('/Title')
                    if title:
                        titles.append(title)
            return titles

        titles = extract_titles(outline)
        print("Merged PDF Bookmarks:", titles)
        
        # Check that top-level file names and inner bookmarks are preserved
        self.assertIn("doc1", titles)
        self.assertIn("Глава 1: Введение", titles)
        self.assertIn("Глава 2: Описание", titles)
        self.assertIn("doc2", titles)
        self.assertIn("Раздел А: Данные", titles)
        self.assertIn("photo", titles)
        self.assertIn("report", titles)

    def test_multipage_tiff_conversion(self):
        # Verify multi-page TIFF frames are all preserved in the converted PDF
        tiff_path = os.path.join(self.test_dir, "multipage.tiff")
        im1 = Image.new("RGB", (200, 200), color=(255, 0, 0))
        im2 = Image.new("RGB", (200, 200), color=(0, 255, 0))
        im1.save(tiff_path, save_all=True, append_images=[im2])
        
        pdf_tmp = convert_image_to_pdf(tiff_path)
        try:
            self.assertTrue(os.path.exists(pdf_tmp))
            r = pypdf.PdfReader(pdf_tmp)
            self.assertEqual(len(r.pages), 2)
        finally:
            if os.path.exists(pdf_tmp):
                os.remove(pdf_tmp)

    def test_locked_output_file_error(self):
        # Verify friendly MergeError is raised if output file is locked by another process
        out_pdf = os.path.join(self.test_dir, "locked_output.pdf")
        with open(out_pdf, "wb") as f_lock:
            f_lock.write(b"%PDF-1.4 dummy")
            # Keep f_lock open to simulate Adobe Acrobat exclusive lock on Windows
            txt_path = os.path.join(self.test_dir, "sample.txt")
            with open(txt_path, "w", encoding="utf-8") as ft:
                ft.write("Hello World")
                
            from src.merger import MergeError
            with self.assertRaises(MergeError) as ctx:
                merge_files([txt_path], out_pdf)
            self.assertIn("в другой программе", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
